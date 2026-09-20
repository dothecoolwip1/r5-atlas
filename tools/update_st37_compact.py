#!/usr/bin/env python3
from __future__ import annotations

import gzip
import json
import math
import os
import re
import shutil
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import requests
from openpyxl import load_workbook

SOURCE_URL = "https://www.aer.ca/prd/documents/sts/st37/ST_37_Excel.zip"
OUT_DIR = Path("testing/data")
META_PATH = OUT_DIR / "st37-meta.json"
SURFACE_PATH = OUT_DIR / "st37-surface.txt.gz"
BORE_PATH = OUT_DIR / "st37-bore.txt.gz"
SEP = "\x1f"


def norm(value):
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    return str(value).strip().replace("\n", " ").replace("\r", " ").replace(SEP, " ")


def number(value):
    if value is None or value == "":
        return None
    try:
        x = float(value)
        return x if math.isfinite(x) else None
    except Exception:
        return None


ALIASES = {
    "licence": ["welllicencenumber", "welllicensenumber", "licencenumber", "licensenumber"],
    "licensee": ["licensee", "companybalongname"],
    "status": ["licencestatus", "licensestatus"],
    "status_date": ["licencestatusdate", "licensestatusdate"],
    "category": ["energydevlpmttype", "energydevlmtcategorytype", "energydevelopmentcategorytype", "welltype", "wellpurpose"],
    "surface_dls": ["licencesurfacelocationlabel", "licensesurfacelocationlabel", "surfacelocation"],
    "sh_lat": ["shactuallatitude", "calculatedlatitude"],
    "sh_lng": ["shactuallongitude", "calculatedlongitude"],
    "uwi": ["welluwi", "uwi"],
    "name": ["wellname"],
    "tmd": ["finaltotaldepth"],
    "tvd": ["maxtrueverticaldepth", "maximumtrueverticaldepth"],
    "bh_lat": ["bhlatitude"],
    "bh_lng": ["bhlongitude"],
}


def find_col(headers, key):
    lookup = {norm(v): i for i, v in enumerate(headers)}
    for alias in ALIASES[key]:
        if alias in lookup:
            return lookup[alias]
    return None


def find_header(ws):
    for row_no, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        vals = [cell_text(v) for v in row]
        n = {norm(v) for v in vals if v is not None}
        if ("welllicencenumber" in n or "welllicensenumber" in n) and (
            "shactuallatitude" in n or "bhlatitude" in n or "welluwi" in n
        ):
            return row_no, vals
    return None, None


def row_value(row, idx):
    return row[idx] if idx is not None and idx < len(row) else None


def open_workbooks(zip_path: Path, temp_dir: Path):
    with zipfile.ZipFile(zip_path) as zf:
        members = [n for n in zf.namelist() if n.lower().endswith((".xlsx", ".xlsm")) and not Path(n).name.startswith("~$")]
        if not members:
            raise RuntimeError("AER ST37 ZIP did not contain an Excel workbook.")
        paths = []
        for member in members:
            name = Path(member).name
            target = temp_dir / name
            with zf.open(member) as src, target.open("wb") as dst:
                shutil.copyfileobj(src, dst)
            paths.append(target)
        return paths


def extract(zip_path: Path):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    surface_count = 0
    bore_count = 0
    surface_seen = False
    bore_seen = False

    with tempfile.TemporaryDirectory() as td:
        workbook_paths = open_workbooks(zip_path, Path(td))
        with gzip.open(SURFACE_PATH, "wt", encoding="utf-8", newline="\n", compresslevel=7) as surface_out,              gzip.open(BORE_PATH, "wt", encoding="utf-8", newline="\n", compresslevel=7) as bore_out:
            surface_out.write("# R5 Atlas compact AER ST37 surface holes\n")
            bore_out.write("# R5 Atlas compact AER ST37 bottom holes\n")

            for workbook_path in workbook_paths:
                wb = load_workbook(workbook_path, read_only=True, data_only=True)
                try:
                    for ws in wb.worksheets:
                        header_row, headers = find_header(ws)
                        if not headers:
                            continue

                        sh_lat = find_col(headers, "sh_lat")
                        sh_lng = find_col(headers, "sh_lng")
                        bh_lat = find_col(headers, "bh_lat")
                        uwi = find_col(headers, "uwi")
                        licence = find_col(headers, "licence")

                        if licence is None:
                            continue

                        if sh_lat is not None and sh_lng is not None:
                            surface_seen = True
                            idx = {
                                k: find_col(headers, k)
                                for k in ("licence", "licensee", "status", "status_date", "category", "surface_dls", "sh_lat", "sh_lng")
                            }
                            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                                lat = number(row_value(row, idx["sh_lat"]))
                                lng = number(row_value(row, idx["sh_lng"]))
                                if lat is None or lng is None or not (48 <= lat <= 61 and -121 <= lng <= -109):
                                    continue
                                licence_value = cell_text(row_value(row, idx["licence"]))
                                if not licence_value:
                                    continue
                                bucket = f"{math.floor(lat * 10)}:{math.floor(lng * 10)}"
                                fields = [
                                    bucket,
                                    f"{lat:.7f}",
                                    f"{lng:.7f}",
                                    licence_value,
                                    cell_text(row_value(row, idx["licensee"])),
                                    cell_text(row_value(row, idx["status"])),
                                    cell_text(row_value(row, idx["status_date"])),
                                    cell_text(row_value(row, idx["category"])),
                                    cell_text(row_value(row, idx["surface_dls"])),
                                ]
                                surface_out.write(SEP.join(fields) + "\n")
                                surface_count += 1

                        if uwi is not None and bh_lat is not None:
                            bore_seen = True
                            idx = {
                                k: find_col(headers, k)
                                for k in ("licence", "uwi", "name", "tmd", "tvd", "bh_lat", "bh_lng", "licensee", "status", "status_date", "category")
                            }
                            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                                licence_value = cell_text(row_value(row, idx["licence"]))
                                if not licence_value:
                                    continue
                                uwi_value = cell_text(row_value(row, idx["uwi"]))
                                lat = number(row_value(row, idx["bh_lat"]))
                                lng = number(row_value(row, idx["bh_lng"]))
                                tmd = number(row_value(row, idx["tmd"]))
                                tvd = number(row_value(row, idx["tvd"]))
                                fields = [
                                    licence_value,
                                    uwi_value,
                                    uwi_value,
                                    cell_text(row_value(row, idx["name"])),
                                    "" if tmd is None else str(tmd),
                                    "" if tvd is None else str(tvd),
                                    "" if lat is None else f"{lat:.7f}",
                                    "" if lng is None else f"{lng:.7f}",
                                    cell_text(row_value(row, idx["licensee"])),
                                    cell_text(row_value(row, idx["status"])),
                                    cell_text(row_value(row, idx["status_date"])),
                                    cell_text(row_value(row, idx["category"])),
                                ]
                                bore_out.write(SEP.join(fields) + "\n")
                                bore_count += 1
                finally:
                    wb.close()

    if not surface_seen or surface_count < 1000:
        raise RuntimeError(f"Surface-hole extraction failed or returned too few records: {surface_count}")
    if not bore_seen or bore_count < 1000:
        raise RuntimeError(f"Bottom-hole extraction failed or returned too few records: {bore_count}")
    return surface_count, bore_count


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    previous = {}
    if META_PATH.exists():
        try:
            previous = json.loads(META_PATH.read_text(encoding="utf-8"))
        except Exception:
            previous = {}

    session = requests.Session()
    session.headers["User-Agent"] = "R5-Atlas-ST37-Updater/1.0"
    head = session.head(SOURCE_URL, allow_redirects=True, timeout=60)
    head.raise_for_status()
    last_modified = head.headers.get("Last-Modified", "")
    etag = head.headers.get("ETag", "")
    source_length = head.headers.get("Content-Length", "")
    force = os.environ.get("FORCE_ST37", "").lower() in {"1", "true", "yes"}

    if not force and previous and (
        (etag and etag == previous.get("sourceEtag")) or
        (last_modified and last_modified == previous.get("sourceLastModified"))
    ) and SURFACE_PATH.exists() and BORE_PATH.exists():
        print("AER ST37 source has not changed. Nothing to rebuild.")
        return 0

    with tempfile.TemporaryDirectory() as td:
        zip_path = Path(td) / "st37.zip"
        with session.get(SOURCE_URL, stream=True, timeout=(60, 600)) as response:
            response.raise_for_status()
            with zip_path.open("wb") as out:
                for chunk in response.iter_content(chunk_size=1024 * 1024):
                    if chunk:
                        out.write(chunk)
        surface_count, bore_count = extract(zip_path)

    now = datetime.now(timezone.utc).isoformat()
    meta = {
        "source": "Alberta Energy Regulator ST37",
        "sourceUrl": SOURCE_URL,
        "sourceLastModified": last_modified,
        "sourceEtag": etag,
        "sourceContentLength": source_length,
        "generatedAt": now,
        "surfaceRecords": surface_count,
        "bottomHoleRecords": bore_count,
        "formatVersion": 1,
    }
    META_PATH.write_text(json.dumps(meta, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(meta, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
